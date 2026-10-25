"""
Profesyonel Rapor Sablonlari -- Technical / Executive / Developer Remediation

Bu modul python-docx ile 3 farkli rapor tipi uretir. Hepsi ayni proje/
session/finding verisinden beslenir, sadece hedef kitleye gore
detay seviyesi ve odak degisir:

  - technical   : Tum test sonuclari + tum bulgular, tam detay (pentester/QA)
  - executive   : Ust duzey ozet, sayilar ve genel risk posturu (yonetim)
  - developer   : Sadece bulgular, "Problem / Neden onemli / Nasil duzeltilir"
                  formatinda, gelistiriciye yonelik somut aksiyon

Remediation rehberligi WSTG kategorisi bazinda GENEL, kamuya acik,
savunma amacli best-practice metinleridir (CWE referanslariyla) -- hicbir
exploit/saldiri detayi icermez.
"""

import io
import json
import os
from datetime import datetime

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Test sonucu etiketleri (TestResult.status)
RESULT_STATUS_LABELS = {
    'pending': 'Bekliyor',
    'passed': 'Acik yok',
    'failed': 'Acik bulundu',
    'skipped': 'Uygulanamaz',
}

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WSTG_DATA_PATH = os.path.join(BASE_DIR, 'data', 'wstg-checklist.en.json')
LLM_DATA_PATH = os.path.join(BASE_DIR, 'data', 'llm-security-checklist.en.json')

SEVERITY_COLORS = {
    'critical': RGBColor(0x99, 0x18, 0x18),
    'high': RGBColor(0xC2, 0x41, 0x0C),
    'medium': RGBColor(0xA1, 0x62, 0x07),
    'low': RGBColor(0x15, 0x80, 0x3D),
    'info': RGBColor(0x47, 0x55, 0x69),
}
SEVERITY_ORDER = ['critical', 'high', 'medium', 'low', 'info']

CATEGORY_REMEDIATION = {
    'WSTG-INFO': {
        'summary': 'Bilgi sizintisini en aza indirin: surum/banner bilgilerini gizleyin, dizin listelemeyi kapatin, gereksiz meta veriyi kaldirin.',
        'cwe': 'CWE-200 (Exposure of Sensitive Information)'
    },
    'WSTG-CONF': {
        'summary': 'Sunucu/uygulama yapilandirmasini sikilastirin: varsayilan/yedek dosyalari kaldirin, yonetim arayuzlerini kisitlayin, en az yetki ilkesini uygulayin.',
        'cwe': 'CWE-16 (Configuration)'
    },
    'WSTG-IDNT': {
        'summary': 'Kimlik yonetimi akislarinda (kayit, kullanici adi) numaralandirmayi zorlastirin; hiz sinirlama ve genel hata mesajlari kullanin.',
        'cwe': 'CWE-204 (Observable Response Discrepancy)'
    },
    'WSTG-ATHN': {
        'summary': 'Guclu parola politikasi, MFA, guvenli kimlik bilgisi tasima ve hesap kilitleme mekanizmalari uygulayin.',
        'cwe': 'CWE-287 (Improper Authentication)'
    },
    'WSTG-ATHZ': {
        'summary': 'Her istekte sunucu tarafi yetkilendirme kontrolu yapin; varsayilan olarak reddet, nesne sahipligini dogrulayin (IDOR onleme).',
        'cwe': 'CWE-285 (Improper Authorization)'
    },
    'WSTG-SESS': {
        'summary': 'Secure/HttpOnly/SameSite cookie bayraklari, giriste session ID yenileme, CSRF token ve uygun oturum zaman asimi uygulayin.',
        'cwe': 'CWE-384 (Session Fixation)'
    },
    'WSTG-INPV': {
        'summary': 'Parametreli sorgular (prepared statements), cikti kodlama, allowlist tabanli girdi dogrulama kullanin; WAF ek savunma katmani olarak dusunulmeli.',
        'cwe': 'CWE-20 (Improper Input Validation)'
    },
    'WSTG-ERRH': {
        'summary': 'Kullaniciya genel hata mesajlari gosterin; detayli stack trace/hata bilgisi yalnizca sunucu tarafli loglarda tutulmali.',
        'cwe': 'CWE-209 (Information Exposure Through an Error Message)'
    },
    'WSTG-CRYP': {
        'summary': 'Guncel/guclu TLS surumleri ve sifre paketleri kullanin, kriptografik anahtarlari guvenli yonetin, eski algoritmalari devre disi birakin.',
        'cwe': 'CWE-327 (Use of a Broken or Risky Cryptographic Algorithm)'
    },
    'WSTG-BUSL': {
        'summary': 'Is mantigi kurallarini ve is akisi durum gecislerini sunucu tarafinda dogrulayin; istemciye guvenmeyin.',
        'cwe': 'CWE-840 (Business Logic Errors)'
    },
    'WSTG-CLNT': {
        'summary': 'Content-Security-Policy uygulayin, DOM ciktisini temizleyin (sanitize), postMessage kullanimini kisitlayin.',
        'cwe': 'CWE-79 (Cross-site Scripting)'
    },
    'WSTG-APIT': {
        'summary': 'Her API endpoint\'inde kimlik dogrulama/yetkilendirme zorunlu kilin, hiz sinirlama ve sema dogrulamasi uygulayin.',
        'cwe': 'CWE-285 (Improper Authorization)'
    },
    # --- OWASP Top 10 for LLM Applications 2025 (v2.0) ---
    'LLM-PI': {
        'summary': 'Guvenilmeyen icerigi (kullanici girdisi, dis belgeler, arac ciktilari) sistem talimatlarindan net sekilde ayirin; kritik eylemlerde ek dogrulama/onay katmani kullanin.',
        'cwe': 'OWASP LLM01:2025 (Prompt Injection)'
    },
    'LLM-SID': {
        'summary': 'Cikti filtreleme ve hassas veri kalip taramasi uygulayin; modelin egitim verisini/PII bilgisini geri uretmesini sinirlayan koruma katmanlari ekleyin.',
        'cwe': 'OWASP LLM02:2025 (Sensitive Information Disclosure)'
    },
    'LLM-SC': {
        'summary': 'Model/eklenti kaynaklarini dogrulanmis surumlere sabitleyin, checksum/imza dogrulamasi yapin, tedarikci guvenlik durusunu duzenli gozden gecirin.',
        'cwe': 'OWASP LLM03:2025 (Supply Chain)'
    },
    'LLM-DMP': {
        'summary': 'Egitim/fine-tuning ve RAG veri kaynaklarina erisim kontrolu ve onay adimlari ekleyin; supheli katkilar icin anomali tespiti kurun.',
        'cwe': 'OWASP LLM04:2025 (Data and Model Poisoning)'
    },
    'LLM-OH': {
        'summary': 'LLM ciktisini asla dogrudan HTML/JS/SQL/shell baglaminda calistirmayin; her zaman cikti kodlama (output encoding) ve sanitization uygulayin.',
        'cwe': 'OWASP LLM05:2025 (Improper Output Handling)'
    },
    'LLM-EA': {
        'summary': 'Ajana taninan arac/izin kapsamini en az yetki ilkesiyle sinirlayin; geri alinamaz/yuksek etkili eylemlerde insan onayi zorunlu kilin.',
        'cwe': 'OWASP LLM06:2025 (Excessive Agency)'
    },
    'LLM-SPL': {
        'summary': 'Sistem promptuna gizli anahtar/kimlik bilgisi veya hassas is mantigi gommeyin; bunlari guvenli backend mantigina tasiyin.',
        'cwe': 'OWASP LLM07:2025 (System Prompt Leakage)'
    },
    'LLM-VEW': {
        'summary': 'Vektor deposuna kaynak veriyle esdeger erisim kontrolu uygulayin; cok kiracili sistemlerde izolasyonu veri erisim katmaninda zorlayin.',
        'cwe': 'OWASP LLM08:2025 (Vector and Embedding Weaknesses)'
    },
    'LLM-MIS': {
        'summary': 'Yuksek riskli kararlarda insan onayli inceleme ve guven/belirsizlik gostergeleri ekleyin; modelin kaynak/alinti uydurmasina karsi dogrulama katmani kurun.',
        'cwe': 'OWASP LLM09:2025 (Misinformation)'
    },
    'LLM-UC': {
        'summary': 'Kullanici/IP bazinda hiz sinirlama ve token/context limitlerini sunucu tarafinda zorunlu kilin; anormal kullanim icin maliyet uyarisi kurun.',
        'cwe': 'OWASP LLM10:2025 (Unbounded Consumption)'
    },
}


def _load_test_index(path):
    """test_id -> {'title':..., 'category_code':..., 'category_name':...}"""
    index = {}
    try:
        with open(path, 'r', encoding='utf-8') as f:
            data = json.load(f)
        for cat in data.get('categories', []):
            for test in cat.get('tests', []):
                index[test['id']] = {
                    'title': test.get('title', ''),
                    'category_code': cat.get('code', ''),
                    'category_name': cat.get('name', '')
                }
    except Exception:
        pass
    return index


# WSTG (web) ve LLM Security (AI) checklist'lerinin ikisi de birlestiriliyor,
# boylece rapor uretimi hangi framework'ten geldigine bakmaksizin (WSTG-* veya
# LLM-*) her bulgu icin dogru baslik/kategori adini gosterebiliyor.
WSTG_INDEX = _load_test_index(WSTG_DATA_PATH)
WSTG_INDEX.update(_load_test_index(LLM_DATA_PATH))


def _category_prefix(test_id):
    parts = test_id.split('-')
    return '-'.join(parts[:2]) if len(parts) >= 2 else test_id


def _set_cell_shading(cell, hex_color):
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _add_page_number_footer(doc):
    section = doc.sections[0]
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_begin = OxmlElement('w:fldChar')
    fld_begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'PAGE'
    fld_end = OxmlElement('w:fldChar')
    fld_end.set(qn('w:fldCharType'), 'end')
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def _cover_page(doc, title_text, project, subtitle=None):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(title_text)
    title_run.bold = True
    title_run.font.size = Pt(28)

    if subtitle:
        sub = doc.add_paragraph()
        sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
        sub_run = sub.add_run(subtitle)
        sub_run.font.size = Pt(14)
        sub_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        "%s\n%s\nOlusturulma: %s" % (
            project.get('name', ''),
            project.get('client', '') or '',
            datetime.utcnow().strftime('%d.%m.%Y %H:%M UTC')
        )
    )
    meta_run.font.size = Pt(12)
    doc.add_page_break()


def _add_severity_table(doc, severity_counts):
    table = doc.add_table(rows=1, cols=len(SEVERITY_ORDER))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, sev in enumerate(SEVERITY_ORDER):
        hdr[i].text = sev.upper()
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        color = SEVERITY_COLORS[sev]
        _set_cell_shading(hdr[i], '%02X%02X%02X' % (color[0], color[1], color[2]))
    row = table.add_row().cells
    for i, sev in enumerate(SEVERITY_ORDER):
        row[i].text = str(severity_counts.get(sev, 0))
        for p in row[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return table


def _severity_counts(findings):
    counts = {k: 0 for k in SEVERITY_ORDER}
    for f in findings:
        sev = (f.get('severity') or 'info').lower()
        if sev in counts:
            counts[sev] += 1
    return counts


def _sorted_findings(findings):
    rank = {s: i for i, s in enumerate(SEVERITY_ORDER)}
    return sorted(findings, key=lambda f: rank.get((f.get('severity') or 'info').lower(), 99))


FINDING_STATUS_LABELS = {
    'open': 'Acik', 'confirmed': 'Dogrulandi', 'fixed': 'Duzeltildi',
    'retest_pending': 'Yeniden test bekliyor', 'retesting': 'Yeniden test ediliyor',
    'resolved': 'Kapatildi', 'wont_fix': 'Duzeltilmeyecek', 'accepted_risk': 'Kabul edilen risk',
}

WSTG_TOTAL = len([k for k in WSTG_INDEX if k.startswith('WSTG-')])


def normalize_finding(f):
    """Bulgular sayfasindaki kayit (models.Finding.to_dict) -> rapor ogesi."""
    return {
        'code': f.get('finding_code') or '',
        'test_id': f.get('test_id') or '',
        'title': f.get('title') or '',
        'severity': (f.get('severity') or 'info').lower(),
        'status': f.get('status') or 'open',
        'description': f.get('description') or '',
        'steps': f.get('steps_to_reproduce') or '',
        'impact': f.get('impact') or '',
        'remediation': f.get('remediation') or '',
        'endpoint': f.get('endpoint') or '',
        'parameter': f.get('parameter') or '',
        'cwe': f.get('cwe') or '',
        'cvss_score': f.get('cvss_score'),
        'cvss_vector': f.get('cvss_vector') or '',
        'references': f.get('references') or '',
    }


def normalize_note(r):
    """Checklist notu (models.TestResult.to_dict) -> rapor ogesi. Yalnizca
    projede Bulgular sayfasinda hic kayit yoksa (eski projeler) kullanilir."""
    meta = WSTG_INDEX.get(r.get('test_id'), {})
    return {
        'code': '', 'test_id': r.get('test_id') or '', 'title': meta.get('title', ''),
        'severity': (r.get('severity') or 'info').lower(), 'status': r.get('finding_status') or 'open',
        'description': r.get('finding') or '', 'steps': '', 'impact': '', 'remediation': '',
        'endpoint': '', 'parameter': '', 'cwe': '', 'cvss_score': None, 'cvss_vector': '', 'references': '',
    }


def _finding_title(f):
    if f.get('code'):
        return "%s - %s" % (f['code'], f['title'])
    if f.get('title'):
        return "%s - %s" % (f['test_id'], f['title']) if f.get('test_id') else f['title']
    return f.get('test_id') or '-'


def _status_label(status):
    return FINDING_STATUS_LABELS.get(status, status or '-')


def _remediation_text(f):
    """Bulgunun kendi cozum onerisi; yoksa WSTG kategorisinin genel onerisi."""
    if f.get('remediation'):
        return f['remediation'], None
    rem = CATEGORY_REMEDIATION.get(_category_prefix(f.get('test_id') or ''))
    if rem:
        return rem['summary'], rem['cwe']
    return '', None


def _add_multiline(doc, text):
    for line in (text or '').strip().splitlines():
        doc.add_paragraph(line)


def _completed_count(all_results):
    return len({r.get('test_id') for r in all_results
                if r.get('status') and r.get('status') != 'pending' and (r.get('test_id') or '').startswith('WSTG-')})


def _source_note(doc, source):
    if source == 'notes':
        p = doc.add_paragraph()
        p.add_run('Not: Bu projede Bulgular sayfasinda kayit olmadigi icin checklist notlari '
                  'bulgu olarak kullanilmistir.').italic = True


def build_technical_report(project, sessions, findings, all_results, source='findings'):
    doc = Document()
    _cover_page(doc, 'Technical Penetration Test Report', project, subtitle='Detayli Teknik Rapor')
    _add_page_number_footer(doc)
    completed = _completed_count(all_results)

    doc.add_heading('1. Proje Bilgileri', level=1)
    info_table = doc.add_table(rows=0, cols=2)
    for label, value in [
        ('Proje', project.get('name', '')),
        ('Musteri', project.get('client', '') or '-'),
        ('Durum', project.get('status', '')),
        ('Baslangic', project.get('start_date') or '-'),
        ('Bitis', project.get('end_date') or '-'),
    ]:
        row = info_table.add_row().cells
        row[0].text = label
        row[0].paragraphs[0].runs[0].bold = True
        row[1].text = str(value)

    doc.add_heading('2. Yonetici Ozeti', level=1)
    doc.add_paragraph(
        "Bu degerlendirme kapsaminda %d/%d WSTG testi tamamlanmis, %d bulgu kaydedilmistir." % (
            completed, WSTG_TOTAL, len(findings)))
    _add_severity_table(doc, _severity_counts(findings))
    doc.add_paragraph()

    doc.add_heading('3. Test Kapsami', level=1)
    by_status = {}
    for r in all_results:
        by_status[r.get('status')] = by_status.get(r.get('status'), 0) + 1
    doc.add_paragraph("WSTG test sayisi: %d" % WSTG_TOTAL)
    doc.add_paragraph("Tamamlanan test: %d" % completed)
    for key in ('passed', 'failed', 'skipped'):
        doc.add_paragraph("%s: %d" % (RESULT_STATUS_LABELS[key], by_status.get(key, 0)), style='List Bullet')

    doc.add_heading('4. Detayli Bulgular', level=1)
    _source_note(doc, source)
    if not findings:
        doc.add_paragraph('Bu degerlendirmede bulgu tespit edilmemistir.')
    for f in _sorted_findings(findings):
        sev = f['severity']
        h = doc.add_heading(_finding_title(f), level=2)
        for run in h.runs:
            run.font.color.rgb = SEVERITY_COLORS.get(sev, SEVERITY_COLORS['info'])

        meta = WSTG_INDEX.get(f['test_id'], {})
        bits = ["Onem: %s" % sev.upper()]
        if f.get('cvss_score') is not None:
            bits.append("CVSS %s" % f['cvss_score'])
        if f.get('test_id'):
            bits.append("Test: %s" % f['test_id'])
        if meta.get('category_name'):
            bits.append("Kategori: %s" % meta['category_name'])
        if f.get('cwe'):
            bits.append(f['cwe'])
        bits.append("Durum: %s" % _status_label(f['status']))
        doc.add_paragraph().add_run('   |   '.join(bits)).italic = True
        if f.get('cvss_vector'):
            doc.add_paragraph().add_run(f['cvss_vector']).italic = True
        if f.get('endpoint'):
            p = doc.add_paragraph()
            p.add_run('Konum: ').bold = True
            p.add_run(f['endpoint'] + (" (parametre: %s)" % f['parameter'] if f.get('parameter') else ''))

        doc.add_paragraph('Bulgu Detayi:').runs[0].bold = True
        _add_multiline(doc, f['description'] or '-')
        if f.get('steps'):
            doc.add_paragraph('Tekrarlama Adimlari:').runs[0].bold = True
            _add_multiline(doc, f['steps'])
        if f.get('impact'):
            doc.add_paragraph('Etki:').runs[0].bold = True
            _add_multiline(doc, f['impact'])
        rem_text, rem_cwe = _remediation_text(f)
        if rem_text:
            doc.add_paragraph('Onerilen Aksiyon:').runs[0].bold = True
            _add_multiline(doc, rem_text)
            if rem_cwe:
                doc.add_paragraph().add_run("Referans: %s" % rem_cwe).italic = True
        if f.get('references'):
            doc.add_paragraph('Referanslar:').runs[0].bold = True
            _add_multiline(doc, f['references'])

    doc.add_heading('5. Tamamlanan Tum Testler', level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Light Grid Accent 1'
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = 'WSTG ID', 'Sonuc', 'Test'
    for r in sorted([r for r in all_results if r.get('status') != 'pending'], key=lambda r: r.get('test_id') or ''):
        row = table.add_row().cells
        row[0].text = r.get('test_id', '')
        row[1].text = RESULT_STATUS_LABELS.get(r.get('status'), r.get('status', ''))
        row[2].text = WSTG_INDEX.get(r.get('test_id'), {}).get('title', '')

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def build_executive_report(project, sessions, findings, all_results, source='findings'):
    doc = Document()
    _cover_page(doc, 'Executive Summary', project, subtitle='Yonetici Ozeti')
    _add_page_number_footer(doc)

    sev_counts = _severity_counts(findings)
    completed = _completed_count(all_results)
    progress = round((completed / WSTG_TOTAL) * 100, 1) if WSTG_TOTAL else 0
    critical_high = sev_counts['critical'] + sev_counts['high']

    doc.add_heading('Genel Risk Posturu', level=1)
    if critical_high > 0:
        risk_text = (
            "Degerlendirme sirasinda %d adet KRITIK/YUKSEK onem dereceli bulgu tespit edilmistir. "
            "Bu bulgularin onceliklendirilerek giderilmesi onerilir." % critical_high
        )
    elif sev_counts['medium'] > 0:
        risk_text = "Kritik veya yuksek seviyede bulgu tespit edilmemistir; orta seviyeli bulgular giderilmelidir."
    else:
        risk_text = "Degerlendirme kapsaminda onemli bir guvenlik bulgusu tespit edilmemistir."
    doc.add_paragraph(risk_text)

    doc.add_paragraph()
    _add_severity_table(doc, sev_counts)
    doc.add_paragraph()

    doc.add_heading('Test Kapsami Ozeti', level=1)
    p = doc.add_paragraph()
    p.add_run("WSTG Kapsami: ").bold = True
    p.add_run("%d/%d test (%%%s)" % (completed, WSTG_TOTAL, progress))
    p2 = doc.add_paragraph()
    p2.add_run("Toplam Bulgu: ").bold = True
    p2.add_run(str(len(findings)))

    doc.add_heading('One Cikan Bulgular', level=1)
    _source_note(doc, source)
    top = [f for f in _sorted_findings(findings) if f['severity'] in ('critical', 'high')][:5]
    if not top:
        doc.add_paragraph('Kritik/Yuksek seviyeli bulgu bulunmamaktadir.')
    for f in top:
        bullet = doc.add_paragraph(style='List Bullet')
        bullet.add_run("%s " % _finding_title(f)).bold = True
        bullet.add_run("(%s, %s)" % (f['severity'].upper(), _status_label(f['status'])))

    doc.add_heading('Sonuc', level=1)
    doc.add_paragraph(
        'Bu rapor, teknik ekiplerin kullanimina yonelik ayrintili Technical Report ile birlikte '
        'degerlendirilmelidir. Gelistirici odakli somut aksiyon adimlari icin Developer Remediation '
        'Report\'a basvurulabilir.'
    )

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def build_developer_report(project, sessions, findings, all_results, source='findings'):
    doc = Document()
    _cover_page(doc, 'Developer Remediation Report', project, subtitle='Gelistirici Aksiyon Rehberi')
    _add_page_number_footer(doc)

    doc.add_paragraph(
        'Bu rapor, tespit edilen her bulgu icin "Problem", "Neden Onemli" ve "Nasil Duzeltilir" '
        'formatinda somut, uygulanabilir rehberlik saglar. Oncelik sirasi severity\'e goredir.'
    )
    _source_note(doc, source)
    doc.add_paragraph()

    # Kapatilmis / kabul edilmis bulgular gelistiriciden aksiyon beklemez.
    actionable = [f for f in findings if f['status'] not in ('resolved', 'accepted_risk', 'wont_fix')]
    if not actionable:
        doc.add_paragraph('Giderilmesi gereken bir bulgu bulunmamaktadir.')

    importance_map = {
        'critical': 'Bu seviyedeki bir zafiyet, saldirganlarin sistemi tamamen ele gecirmesine veya kritik veriye erismesine yol acabilir. Acil giderilmelidir.',
        'high': 'Bu zafiyet ciddi veri sizintisi veya yetkisiz erisime yol acabilir. Yuksek oncelikle giderilmelidir.',
        'medium': 'Bu zafiyet tek basina sinirli etkiye sahip olsa da baska zafiyetlerle birlestiginde risk olusturabilir.',
        'low': 'Dogrudan risk dusuk olsa da savunma derinligi ilkesi geregi giderilmesi onerilir.',
        'info': 'Bilgilendirme amaclidir, dogrudan bir risk teskil etmeyebilir ancak gozden gecirilmelidir.',
    }

    for f in _sorted_findings(actionable):
        sev = f['severity']
        h = doc.add_heading(_finding_title(f), level=2)
        for run in h.runs:
            run.font.color.rgb = SEVERITY_COLORS.get(sev, SEVERITY_COLORS['info'])

        badge_p = doc.add_paragraph()
        badge_p.add_run("Onem: %s" % sev.upper()).bold = True
        where = f.get('endpoint') or f.get('test_id') or '-'
        badge_p.add_run("   |   Konum: %s   |   Durum: %s" % (where, _status_label(f['status'])))

        doc.add_paragraph('Problem:').runs[0].bold = True
        _add_multiline(doc, f['description'] or '-')
        if f.get('steps'):
            doc.add_paragraph('Nasil Tekrarlanir:').runs[0].bold = True
            _add_multiline(doc, f['steps'])

        doc.add_paragraph('Neden Onemli:').runs[0].bold = True
        _add_multiline(doc, f.get('impact') or importance_map.get(sev, ''))

        doc.add_paragraph('Nasil Duzeltilir:').runs[0].bold = True
        rem_text, rem_cwe = _remediation_text(f)
        if rem_text:
            _add_multiline(doc, rem_text)
            refs = [x for x in (rem_cwe or f.get('cwe'), f.get('test_id')) if x]
            if refs:
                doc.add_paragraph().add_run("Referans: %s" % '  |  '.join(refs)).italic = True
        else:
            doc.add_paragraph('Ilgili WSTG kategorisinin best-practice rehberine basvurun.')

        doc.add_paragraph()

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


REPORT_BUILDERS = {
    'technical': build_technical_report,
    'executive': build_executive_report,
    'developer': build_developer_report,
}

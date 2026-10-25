"""Word raporları (teknik / yönetici / geliştirici) Bulgular sayfasındaki kayıtları kullanır."""
import io
import os
import re
import sys

sys.path.insert(0, '.')
os.environ['FLASK_ENV'] = 'development'

import docx
import pytest

import app as app_module
import report_builder


@pytest.fixture
def client():
    app_module.app.config['TESTING'] = True
    return app_module.app.test_client()


def _text(resp):
    assert resp.status_code == 200, resp.get_data(as_text=True)
    d = docx.Document(io.BytesIO(resp.data))
    parts = [p.text for p in d.paragraphs]
    parts += [c.text for t in d.tables for row in t.rows for c in row.cells]
    return '\n'.join(parts)


@pytest.fixture
def project(client):
    pid = client.post('/api/projects', json={'name': 'Word Rapor', 'client': 'Acme'}).get_json()['id']
    sid = client.post('/api/sessions', json={'name': 'S', 'project_id': pid}).get_json()['id']
    # Checklist notu: Bulgular sayfasında kayıt varken rapora girmemeli
    client.post(f'/api/sessions/{sid}/results', json={'test_id': 'WSTG-INFO-01', 'status': 'passed',
                                                     'finding': 'CHECKLIST-NOTU-METNI', 'severity': 'low'})
    client.post(f'/api/sessions/{sid}/results', json={'test_id': 'WSTG-INPV-01', 'status': 'failed'})
    client.post(f'/api/sessions/{sid}/results', json={'test_id': 'WSTG-CRYP-01', 'status': 'skipped'})
    client.post(f'/api/projects/{pid}/findings', json={
        'title': 'Arama kutusunda yansıtılan XSS', 'severity': 'high', 'test_id': 'WSTG-INPV-01',
        'endpoint': '/search', 'parameter': 'q', 'description': 'q parametresi kaçışsız yansıtılıyor.',
        'steps_to_reproduce': '1. /search?q=<svg/onload=alert(1)> aç\n2. Uyarı kutusunu gör',
        'impact': 'Oturum çerezi çalınabilir.', 'remediation': 'Çıktıyı bağlama göre kodlayın.', 'cwe': 'CWE-79'})
    client.post(f'/api/projects/{pid}/findings', json={'title': 'Eski TLS sürümü', 'severity': 'low', 'status': 'resolved',
                                                     'test_id': 'WSTG-CRYP-01'})
    return pid, sid


def test_technical_report_uses_findings(client, project):
    pid, _ = project
    text = _text(client.get(f'/api/projects/{pid}/reports/technical'))
    assert re.search(r'FND-\d{4} - Arama kutusunda yansıtılan XSS', text)
    assert 'Tekrarlama Adimlari:' in text and '/search?q=<svg/onload=alert(1)> aç' in text
    assert 'Çıktıyı bağlama göre kodlayın.' in text and 'Oturum çerezi çalınabilir.' in text
    assert 'Konum: /search (parametre: q)' in text and 'CWE-79' in text
    assert 'Kapatildi' in text                                   # durum etiketi okunur
    assert 'CHECKLIST-NOTU-METNI' not in text
    assert '3/%d WSTG testi tamamlanmis, 2 bulgu' % report_builder.WSTG_TOTAL in text
    assert 'Acik bulundu: 1' in text and 'Uygulanamaz: 1' in text and 'Acik yok: 1' in text


def test_developer_report_skips_closed_and_shows_steps(client, project):
    pid, _ = project
    text = _text(client.get(f'/api/projects/{pid}/reports/developer'))
    assert 'Arama kutusunda yansıtılan XSS' in text and 'Nasil Tekrarlanir:' in text
    assert 'Eski TLS sürümü' not in text                        # kapatılmış: aksiyon gerekmez
    assert 'Oturum çerezi çalınabilir.' in text                 # "Neden önemli" bulgunun etkisinden


def test_executive_report_coverage_and_top_findings(client, project):
    pid, _ = project
    text = _text(client.get(f'/api/projects/{pid}/reports/executive'))
    assert '1 adet KRITIK/YUKSEK' in text
    assert '3/%d test' % report_builder.WSTG_TOTAL in text
    assert re.search(r'FND-\d{4} - Arama kutusunda yansıtılan XSS \(HIGH, Acik\)', text)


def test_old_projects_fall_back_to_checklist_notes(client):
    pid = client.post('/api/projects', json={'name': 'Eski'}).get_json()['id']
    sid = client.post('/api/sessions', json={'name': 'S', 'project_id': pid}).get_json()['id']
    client.post(f'/api/sessions/{sid}/results', json={'test_id': 'WSTG-ATHN-01', 'status': 'failed',
                                                     'finding': 'Giriş HTTP üzerinden.', 'severity': 'high'})
    text = _text(client.get(f'/api/projects/{pid}/reports/technical'))
    assert 'Giriş HTTP üzerinden.' in text and 'checklist notlari' in text


def test_empty_project_reports(client):
    pid = client.post('/api/projects', json={'name': 'Boş'}).get_json()['id']
    for kind in ('technical', 'executive', 'developer'):
        _text(client.get(f'/api/projects/{pid}/reports/{kind}'))
